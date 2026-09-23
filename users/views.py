from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth import login, logout
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.db.models import Sum, Q
from django.utils import timezone
import json
from .forms import CustomUserCreationForm, CustomLoginForm
from .models import CustomUser, Role
from .decorators import role_required
from books.models import Book, BookCopy
from circulation.models import Loan, Fine, Reservation
from django.db.models import Count
from django.db.models.functions import TruncMonth
from datetime import timedelta
from django.utils import timezone
from books.models import Book, BookCopy
from circulation.models import Loan, Fine, Reservation
from users.models import CustomUser
from django.db.models import Q, Count, Sum
from django.http import HttpResponse, FileResponse
from io import BytesIO
from openpyxl import Workbook
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from reportlab.lib.units import inch
from .models import Notification
from django.http import JsonResponse

# ==================== HOME ====================
def home(request):
    if request.user.is_authenticated:
        if request.user.is_admin() or request.user.is_librarian():
            return redirect('staff_dashboard')
        return redirect('dashboard')
    return render(request, 'home.html')


# ==================== AUTHENTICATION ====================
def login_view(request):
    if request.user.is_authenticated:
        return redirect('home')

    if request.method == 'POST':
        form = CustomLoginForm(request, data=request.POST)
        if form.is_valid():
            user = form.get_user()
            login(request, user)
            messages.success(request, f"Welcome back, {user.get_full_name() or user.username}!")

            if user.is_admin() or user.is_librarian():
                return redirect('staff_dashboard')
            return redirect('dashboard')
        else:
            messages.error(request, "Invalid username or password.")
    else:
        form = CustomLoginForm()

    return render(request, 'users/login.html', {'form': form})


def register_view(request):
    if request.user.is_authenticated:
        return redirect('home')

    if request.method == 'POST':
        form = CustomUserCreationForm(request.POST)
        if form.is_valid():
            user = form.save(commit=False)
            reader_role = Role.objects.get(role_name='Reader')
            user.role = reader_role
            user.save()
            messages.success(request, "Account created successfully! Please login.")
            return redirect('login')
        else:
            messages.error(request, "Please correct the errors below.")
    else:
        form = CustomUserCreationForm()

    return render(request, 'users/register.html', {'form': form})


def logout_view(request):
    logout(request)
    messages.info(request, "You have been logged out successfully.")
    return redirect('login')


# ==================== READER DASHBOARD ====================
@login_required
@role_required('Reader')
def reader_dashboard(request):
    user = request.user
    active_loans = Loan.objects.filter(user=user, status='Active').select_related('copy__book')
    overdue_loans = active_loans.filter(due_date__lt=timezone.now().date())
    reservations = Reservation.objects.filter(user=user, status='Pending').select_related('book')



    # ========== Chart Data ==========
    # Books issued in last 6 months
    six_months_ago = timezone.now().date() - timedelta(days=180)
    monthly_issues = (
        Loan.objects.filter(loan_date__gte=six_months_ago)
        .annotate(month=TruncMonth('loan_date'))
        .values('month')
        .annotate(count=Count('id'))
        .order_by('month')
    )

    chart_labels = [item['month'].strftime('%b %Y') for item in monthly_issues]
    chart_data = [item['count'] for item in monthly_issues]

    # Most borrowed books (Top 5)
    top_books = (
        Loan.objects.values('copy__book__title')
        .annotate(total=Count('id'))
        .order_by('-total')[:5]
    )
    top_book_labels = [item['copy__book__title'] or 'Unknown' for item in top_books]
    top_book_data = [item['total'] for item in top_books]
    

    context = {
        'active_loans': active_loans,
        'overdue_count': overdue_loans.count(),
        'reservation_count': reservations.count(),
        'total_fines': user.total_fines,
        'reservations': reservations,
        'chart_labels': chart_labels,
        'chart_data': chart_data,
        'top_book_labels': top_book_labels,
        'top_book_data': top_book_data,
        
    }
    return render(request, 'users/reader_dashboard.html', context)


# ==================== STAFF DASHBOARD ====================
@login_required
@role_required('Admin', 'Librarian')
def staff_dashboard(request):
    import json
    from django.db.models import Count, Sum
    from django.db.models.functions import TruncMonth
    from datetime import timedelta

    today = timezone.now().date()

    # Statistics
    total_books = Book.objects.count()
    available_books = BookCopy.objects.filter(status='Available').count()
    total_members = CustomUser.objects.filter(role__role_name='Reader').count()
    total_staff = CustomUser.objects.filter(role__role_name__in=['Admin', 'Librarian']).count()
    books_issued_today = Loan.objects.filter(loan_date=today).count()
    books_returned_today = Loan.objects.filter(return_date=today).count()
    overdue_books = Loan.objects.filter(status='Active', due_date__lt=today).count()
    pending_fine = Fine.objects.filter(status__in=['Unpaid', 'Partially Paid']).aggregate(total=Sum('amount'))['total'] or 0
    pending_reservations = Reservation.objects.filter(status='Pending').count()
    recent_loans = Loan.objects.select_related('user', 'copy__book').order_by('-loan_date')[:10]

    # Chart Data
    six_months_ago = today - timedelta(days=180)
    monthly_issues = (
        Loan.objects.filter(loan_date__gte=six_months_ago)
        .annotate(month=TruncMonth('loan_date'))
        .values('month')
        .annotate(count=Count('id'))
        .order_by('month')
    )

    chart_labels = []
    chart_data = []
    for item in monthly_issues:
        if item['month']:
            chart_labels.append(item['month'].strftime('%b %Y'))
            chart_data.append(item['count'])

    top_books = (
        Loan.objects.values('copy__book__title')
        .annotate(total=Count('id'))
        .order_by('-total')[:5]
    )
    top_book_labels = [item['copy__book__title'] or 'Unknown' for item in top_books]
    top_book_data = [item['total'] for item in top_books]

    context = {
        'total_books': total_books,
        'available_books': available_books,
        'total_members': total_members,
        'total_staff': total_staff,
        'books_issued_today': books_issued_today,
        'books_returned_today': books_returned_today,
        'overdue_books': overdue_books,
        'pending_fine': pending_fine,
        'pending_reservations': pending_reservations,
        'recent_loans': recent_loans,
        'chart_labels': json.dumps(chart_labels),
        'chart_data': json.dumps(chart_data),
        'top_book_labels': json.dumps(top_book_labels),
        'top_book_data': json.dumps(top_book_data),
    }
    return render(request, 'users/staff_dashboard.html', context)


# ==================== PHASE 5: READER MANAGEMENT ====================
@login_required
@role_required('Admin', 'Librarian')
def reader_list(request):
    readers = CustomUser.objects.filter(role__role_name='Reader').order_by('-date_joined')
    query = request.GET.get('q', '')
    if query:
        readers = readers.filter(
            Q(username__icontains=query) |
            Q(first_name__icontains=query) |
            Q(last_name__icontains=query) |
            Q(email__icontains=query)
        )
    return render(request, 'users/reader_list.html', {'readers': readers, 'query': query})


@login_required
@role_required('Admin', 'Librarian')
def create_reader(request):
    if request.method == 'POST':
        form = CustomUserCreationForm(request.POST)
        if form.is_valid():
            user = form.save(commit=False)
            user.role = Role.objects.get(role_name='Reader')
            user.save()
            messages.success(request, f"Reader '{user.username}' created successfully!")
            return redirect('reader_list')
    else:
        form = CustomUserCreationForm()
    return render(request, 'users/create_reader.html', {'form': form})


@login_required
@role_required('Admin', 'Librarian')
def edit_reader(request, user_id):
    reader = get_object_or_404(CustomUser, id=user_id, role__role_name='Reader')
    if request.method == 'POST':
        reader.first_name = request.POST.get('first_name')
        reader.last_name = request.POST.get('last_name')
        reader.email = request.POST.get('email')
        reader.phone = request.POST.get('phone')
        reader.borrowing_limit = request.POST.get('borrowing_limit', 5)
        reader.is_active = request.POST.get('is_active') == 'on'
        reader.save()
        messages.success(request, "Reader updated successfully!")
        return redirect('reader_detail', user_id=reader.id)
    return render(request, 'users/edit_reader.html', {'reader': reader})


@login_required
@role_required('Admin', 'Librarian')
def delete_reader(request, user_id):
    reader = get_object_or_404(CustomUser, id=user_id, role__role_name='Reader')
    if request.method == 'POST':
        username = reader.username
        reader.delete()
        messages.success(request, f"Reader '{username}' deleted successfully!")
        return redirect('reader_list')
    return render(request, 'users/delete_reader.html', {'reader': reader})


@login_required
@role_required('Admin', 'Librarian')
def reader_detail(request, user_id):
    reader = get_object_or_404(CustomUser, id=user_id, role__role_name='Reader')
    loans = Loan.objects.filter(user=reader).select_related('copy__book').order_by('-loan_date')
    fines = Fine.objects.filter(user=reader).order_by('-fine_date')
    return render(request, 'users/reader_detail.html', {
        'reader': reader,
        'loans': loans,
        'fines': fines,
    })  


# ==================== PROFILE ====================
@login_required
def profile(request):
    return render(request, 'users/profile.html', {'user': request.user})


@login_required
def download_library_card(request):
    """
    Generates a printable PDF membership card for the logged-in user.
    (reportlab was already in requirements.txt but nothing used it yet.)
    """
    user = request.user
    buffer = BytesIO()

    card_width, card_height = 4.5 * inch, 2.8 * inch
    p = canvas.Canvas(buffer, pagesize=(card_width, card_height))
    primary = (79 / 255, 70 / 255, 229 / 255)  # matches --primary in style.css

    # Header
    p.setFillColorRGB(*primary)
    p.rect(0, card_height - 0.65 * inch, card_width, 0.65 * inch, fill=1, stroke=0)
    p.setFillColorRGB(1, 1, 1)
    p.setFont('Helvetica-Bold', 16)
    p.drawString(0.25 * inch, card_height - 0.42 * inch, "ONLINE LIBRARY")
    p.setFont('Helvetica', 8)
    p.drawString(0.25 * inch, card_height - 0.58 * inch, "MEMBER IDENTIFICATION CARD")

    # Body
    y = card_height - 0.95 * inch
    p.setFillColorRGB(0.12, 0.16, 0.22)
    p.setFont('Helvetica-Bold', 13)
    p.drawString(0.25 * inch, y, user.get_full_name() or user.username)

    details = [
        f"Member ID: LIB-{user.id:05d}",
        f"Role: {user.role.role_name if user.role else 'Reader'}",
        f"Member Since: {user.date_joined.strftime('%d %b %Y')}",
        f"Valid Until: {user.membership_expiry.strftime('%d %b %Y') if user.membership_expiry else 'No expiry set'}",
        f"Borrowing Limit: {user.borrowing_limit} books",
    ]
    p.setFont('Helvetica', 9)
    for line in details:
        y -= 0.22 * inch
        p.drawString(0.25 * inch, y, line)

    # Footer
    p.setStrokeColorRGB(0.85, 0.85, 0.85)
    p.line(0.25 * inch, 0.3 * inch, card_width - 0.25 * inch, 0.3 * inch)
    p.setFillColorRGB(0.5, 0.5, 0.5)
    p.setFont('Helvetica-Oblique', 7)
    p.drawString(0.25 * inch, 0.15 * inch, "Present this card at the circulation desk. Not transferable.")

    p.showPage()
    p.save()
    buffer.seek(0)
    return FileResponse(buffer, as_attachment=True, filename=f"library_card_{user.username}.pdf")


@login_required
def edit_profile(request):
    user = request.user

    if request.method == 'POST':
        user.first_name = request.POST.get('first_name')
        user.last_name = request.POST.get('last_name')
        user.email = request.POST.get('email')
        user.phone = request.POST.get('phone')
        user.address = request.POST.get('address', '')
        user.save()
        messages.success(request, "Profile updated successfully!")
        return redirect('profile')

    return render(request, 'users/edit_profile.html', {'user': user})


# ==================== FINE HISTORY ====================
@login_required
@role_required('Reader')
def fine_history(request):
    fines = Fine.objects.filter(user=request.user).order_by('-fine_date')
    return render(request, 'users/fine_history.html', {'fines': fines})


@login_required
@role_required('Admin', 'Librarian')
def reports(request):
    import json
    from django.db.models import Count, Sum
    from django.db.models.functions import TruncMonth
    from datetime import timedelta

    today = timezone.now().date()

    # 1. Most Borrowed Books (Top 10)
    most_borrowed = (
        Loan.objects.values('copy__book__title')
        .annotate(total=Count('id'))
        .order_by('-total')[:10]
    )
    most_borrowed_labels = [i['copy__book__title'] or 'Unknown' for i in most_borrowed]
    most_borrowed_data = [i['total'] for i in most_borrowed]

    # 2. Fine Report
    total_fines = Fine.objects.aggregate(total=Sum('amount'))['total'] or 0
    paid_fines = Fine.objects.filter(status='Paid').aggregate(total=Sum('paid_amount'))['total'] or 0
    pending_fines = Fine.objects.filter(status__in=['Unpaid', 'Partially Paid']).aggregate(total=Sum('amount'))['total'] or 0

    # 3. Active Readers
    active_readers = (
        CustomUser.objects.filter(role__role_name='Reader', loans__status='Active')
        .distinct()
        .annotate(active_loans=Count('loans', filter=Q(loans__status='Active')))
        .order_by('-active_loans')[:10]
    )

    # 4. Monthly Rentals (Last 6 months)
    six_months_ago = today - timedelta(days=180)
    monthly = (
        Loan.objects.filter(loan_date__gte=six_months_ago)
        .annotate(month=TruncMonth('loan_date'))
        .values('month')
        .annotate(count=Count('id'))
        .order_by('month')
    )
    monthly_labels = [i['month'].strftime('%b %Y') for i in monthly if i['month']]
    monthly_data = [i['count'] for i in monthly]

    # 5. Category Statistics
    category_stats = (
        Loan.objects.values('copy__book__category__category_name')
        .annotate(total=Count('id'))
        .order_by('-total')
    )
    category_labels = [i['copy__book__category__category_name'] or 'Uncategorized' for i in category_stats]
    category_data = [i['total'] for i in category_stats]

    context = {
        'most_borrowed': most_borrowed,
        'most_borrowed_labels': json.dumps(most_borrowed_labels),
        'most_borrowed_data': json.dumps(most_borrowed_data),
        'total_fines': total_fines,
        'paid_fines': paid_fines,
        'pending_fines': pending_fines,
        'active_readers': active_readers,
        'monthly_labels': json.dumps(monthly_labels),
        'monthly_data': json.dumps(monthly_data),
        'category_labels': json.dumps(category_labels),
        'category_data': json.dumps(category_data),
    }
    return render(request, 'users/reports.html', context)


@login_required
@role_required('Admin', 'Librarian')
def export_most_borrowed_excel(request):
    most_borrowed = (
        Loan.objects.values('copy__book__title')
        .annotate(total=Count('id'))
        .order_by('-total')[:20]
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Most Borrowed Books"
    ws.append(["Rank", "Book Title", "Times Borrowed"])

    for i, item in enumerate(most_borrowed, start=1):
        ws.append([i, item['copy__book__title'] or "Unknown", item['total']])

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename=most_borrowed_books.xlsx'
    wb.save(response)
    return response


@login_required
@role_required('Admin', 'Librarian')
def export_fines_excel(request):
    fines = Fine.objects.select_related('user').order_by('-fine_date')

    wb = Workbook()
    ws = wb.active
    ws.title = "Fine Report"
    ws.append(["User", "Email", "Amount", "Paid", "Status", "Date"])

    for fine in fines:
        ws.append([
            fine.user.get_full_name() or fine.user.username,
            fine.user.email,
            float(fine.amount),
            float(fine.paid_amount or 0),
            fine.status,
            str(fine.fine_date),
        ])

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename=fine_report.xlsx'
    wb.save(response)
    return response


@login_required
@role_required('Admin', 'Librarian')
def export_active_readers_excel(request):
    active_readers = (
        CustomUser.objects.filter(role__role_name='Reader', loans__status='Active')
        .distinct()
        .annotate(active_loans=Count('loans', filter=Q(loans__status='Active')))
        .order_by('-active_loans')
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Active Readers"
    ws.append(["Name", "Email", "Phone", "Active Loans"])

    for reader in active_readers:
        ws.append([
            reader.get_full_name() or reader.username,
            reader.email,
            reader.phone or "",
            reader.active_loans,
        ])

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename=active_readers.xlsx'
    wb.save(response)
    return response



@login_required
def get_notifications(request):
    notifications = Notification.objects.filter(user=request.user)[:10]
    unread_count = Notification.objects.filter(user=request.user, is_read=False).count()

    data = []
    for n in notifications:
        data.append({
            'id': n.id,
            'title': n.title,
            'message': n.message,
            'is_read': n.is_read,
            'created_at': n.created_at.strftime('%d %b %Y %H:%M'),
            'link': n.link or '',
        })

    return JsonResponse({
        'notifications': data,
        'unread_count': unread_count
    })


@login_required
def mark_notification_read(request, notification_id):
    notification = get_object_or_404(Notification, id=notification_id, user=request.user)
    notification.is_read = True
    notification.save()
    return JsonResponse({'success': True})


@login_required
def mark_all_notifications_read(request):
    Notification.objects.filter(user=request.user, is_read=False).update(is_read=True)
    return JsonResponse({'success': True})